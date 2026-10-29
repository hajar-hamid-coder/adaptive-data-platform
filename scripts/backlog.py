#!/usr/bin/env python3
"""Interface unique avec backlog/backlog.xlsx (source de vérité du projet).

Usage :
  backlog.py status                      état par phase, phase courante, prochaine tâche
  backlog.py next                        détail de la prochaine tâche à exécuter
  backlog.py show P1-03                  détail d'une tâche
  backlog.py set P1-03 "En cours" [--note "..."]
  backlog.py validate P1                 gate : À valider -> Validé (après accord de l'utilisateur)
  backlog.py hypo [H-01 ...]             hypothèses ouvertes (ou celles citées)
  backlog.py hypo-add --tasks .. --question .. --assumption .. --impact .. --ask ..
  backlog.py export-md                   régénère backlog/BACKLOG.md
  backlog.py hook-session | hook-stop    sorties pour les hooks Claude Code
"""
import argparse
import datetime
import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl manquant : `uv sync` ou `pip install openpyxl`", file=sys.stderr)
    sys.exit(2)

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "backlog" / "backlog.xlsx"
MD = ROOT / "backlog" / "BACKLOG.md"
STATE = ROOT / ".claude" / "state"
FIRST, SECOND, DONE, VALID = "Non fait", "En cours", "À valider", "Validé"
ALLOWED = {(FIRST, SECOND), (SECOND, DONE), (SECOND, FIRST), (DONE, SECOND)}
SKIP = {"Ordre_Execution", "Hypotheses", "Legende"}


class Backlog:
    def __init__(self):
        self.wb = load_workbook(XLSX)
        self.ctl = self.wb["Ordre_Execution"]
        self.phases = []  # (pid, sheet, title, row)
        for r in range(5, 40):
            pid = self.ctl.cell(r, 1).value
            if isinstance(pid, str) and pid.startswith("P") and self.ctl.cell(r, 2).value in self.wb.sheetnames:
                self.phases.append((pid, self.ctl.cell(r, 2).value, self.ctl.cell(r, 3).value, r))
        self.tasks = []
        for pid, sheet, _, _ in self.phases:
            ws = self.wb[sheet]
            head = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, 1).value:
                    row = {k: ws.cell(r, c).value for k, c in head.items()}
                    row.update(_phase=pid, _ws=ws, _row=r, _head=head)
                    self.tasks.append(row)

    def gate(self, pid):
        for p, _, _, r in self.phases:
            if p == pid:
                return self.ctl.cell(r, 12).value == "Oui"
        raise KeyError(pid)

    def task(self, tid):
        for t in self.tasks:
            if t["ID"] == tid:
                return t
        raise SystemExit(f"Tâche inconnue : {tid}")

    def of_phase(self, pid):
        return [t for t in self.tasks if t["_phase"] == pid]

    def current_phase(self):
        for pid, *_ in self.phases:
            if not self.gate(pid):
                return pid
        return None

    def next_task(self):
        pid = self.current_phase()
        if pid is None:
            return None
        for t in self.of_phase(pid):
            if t["Statut"] in (FIRST, SECOND):
                return t
        return None

    def save(self):
        self.wb.save(XLSX)
        (STATE / "stop_nudges").unlink(missing_ok=True) if STATE.exists() else None


def deps(t):
    raw = (t.get("Dépendances") or "").replace("—", "")
    return [d.strip() for d in raw.split(",") if d.strip()]


def fmt(t):
    keys = ["ID", "Tâche", "Objectif", "Code à produire", "Explication & compétence CV", "Test de validation",
            "Modèle", "Dépendances", "Hypothèse", "Skills / plugins", "Statut", "Preuve / notes Claude"]
    return "\n".join(f"{k}: {t[k]}" for k in keys if t.get(k) not in (None, ""))


def counts(b, pid):
    ts = b.of_phase(pid)
    return {s: sum(1 for t in ts if t["Statut"] == s) for s in (FIRST, SECOND, DONE, VALID)}, len(ts)


def cmd_status(b, hook=False):
    cur = b.current_phase()
    for pid, _, title, _ in b.phases:
        c, n = counts(b, pid)
        mark = "▶" if pid == cur else ("✔" if b.gate(pid) else " ")
        print(f"{mark} {pid} {title}: {c[VALID]}/{n} validées, {c[DONE]} à valider, {c[SECOND]} en cours, {c[FIRST]} à faire")
    if cur is None:
        print("Toutes les phases sont validées.")
        return
    t = b.next_task()
    if t:
        print(f"Prochaine tâche : {t['ID']} — {t['Tâche']} (modèle : {t['Modèle']}, statut : {t['Statut']})")
    else:
        print(f"Phase {cur} : toutes les tâches sont 'À valider'. ACTION : présenter le bilan et demander la validation à l'utilisateur.")


def cmd_set(b, tid, statut, note):
    t = b.task(tid)
    if statut == VALID:
        raise SystemExit("Interdit : 'Validé' n'est posé que par `backlog.py validate <phase>` après accord de l'utilisateur.")
    if statut not in (FIRST, SECOND, DONE):
        raise SystemExit(f"Statut invalide : {statut}")
    if (t["Statut"], statut) not in ALLOWED and t["Statut"] != statut:
        raise SystemExit(f"Transition refusée : {t['Statut']} -> {statut}")
    pids = [p for p, *_ in b.phases]
    idx = pids.index(t["_phase"])
    if idx and not b.gate(pids[idx - 1]):
        raise SystemExit(f"La phase {pids[idx - 1]} n'est pas validée par l'utilisateur : phase {t['_phase']} verrouillée.")
    if statut == SECOND:
        for d in deps(t):
            if b.task(d)["Statut"] not in (DONE, VALID):
                raise SystemExit(f"Dépendance non terminée : {d} ({b.task(d)['Statut']})")
        busy = [x["ID"] for x in b.tasks if x["Statut"] == SECOND and x["ID"] != tid]
        if busy:
            raise SystemExit(f"Une seule tâche 'En cours' à la fois (déjà : {', '.join(busy)})")
    ws, r, h = t["_ws"], t["_row"], t["_head"]
    ws.cell(r, h["Statut"], statut)
    if note:
        stamp = datetime.date.today().isoformat()
        old = ws.cell(r, h["Preuve / notes Claude"]).value
        ws.cell(r, h["Preuve / notes Claude"], (f"{old}\n" if old else "") + f"[{stamp}] {statut}: {note}")
    b.save()
    print(f"{tid} -> {statut}")


def cmd_validate(b, pid):
    ts = b.of_phase(pid)
    if not ts:
        raise SystemExit(f"Phase inconnue : {pid}")
    bad = [t["ID"] for t in ts if t["Statut"] not in (DONE, VALID)]
    if bad:
        raise SystemExit(f"Validation impossible, tâches non terminées : {', '.join(bad)}")
    for t in ts:
        t["_ws"].cell(t["_row"], t["_head"]["Statut"], VALID)
    for p, _, _, r in b.phases:
        if p == pid:
            b.ctl.cell(r, 12, "Oui")
    b.save()
    print(f"Phase {pid} validée ({len(ts)} tâches).")


def cmd_export_md(b):
    out = ["# Backlog — Plateforme de Données Intelligente et Évolutive", "",
           "> Fichier généré depuis `backlog.xlsx` (source de vérité). Ne pas éditer : `python scripts/backlog.py export-md`.", ""]
    for pid, sheet, title, r in b.phases:
        goal = b.ctl.cell(r, 4).value
        out += [f"## {pid} — {title}", "", f"*{goal}* — gate utilisateur : **{b.ctl.cell(r, 12).value}**", "",
                "| Ordre | ID | Tâche | Modèle | Dépend. | Statut |", "|---|---|---|---|---|---|"]
        for t in b.of_phase(pid):
            out.append(f"| {t['Ordre']} | {t['ID']} | {t['Tâche']} | {t['Modèle']} | {t['Dépendances']} | {t['Statut']} |")
        out.append("")
        for t in b.of_phase(pid):
            out += [f"### {t['ID']} — {t['Tâche']}", f"- **Objectif** : {t['Objectif']}", f"- **Code** : {t['Code à produire']}",
                    f"- **Pourquoi** : {t['Explication & compétence CV']}", f"- **Test** : {t['Test de validation']}",
                    f"- **Hypothèse** : {t['Hypothèse']} · **Skills** : {t['Skills / plugins']}", ""]
    MD.write_text("\n".join(out), encoding="utf-8")
    print(f"écrit {MD.relative_to(ROOT)}")


def cmd_hypo(b, ids):
    ws = b.wb["Hypotheses"]
    for r in range(2, ws.max_row + 1):
        hid, status = ws.cell(r, 1).value, ws.cell(r, 7).value
        if not hid or (ids and hid not in ids) or (not ids and status != "Ouverte"):
            continue
        print(f"{hid} [{status}] tâches: {ws.cell(r, 2).value}\n  Q: {ws.cell(r, 3).value}\n  Hypothèse: {ws.cell(r, 4).value}\n"
              f"  Impact: {ws.cell(r, 5).value}\n  À demander: {ws.cell(r, 6).value}\n  Décision: {ws.cell(r, 8).value or '-'}")


def cmd_hypo_add(b, a):
    ws = b.wb["Hypotheses"]
    r = ws.max_row + 1
    last = max(int(str(ws.cell(i, 1).value).split("-")[1]) for i in range(2, ws.max_row + 1))
    row = [f"H-{last + 1:02d}", a.tasks, a.question, a.assumption, a.impact, a.ask, "Ouverte", "", ""]
    for c, v in enumerate(row, 1):
        ws.cell(r, c, v)._style = ws.cell(r - 1, c)._style
    b.save()
    print(f"{row[0]} ajoutée")


def cmd_hook_session(b):
    print("=== État du backlog (scripts/backlog.py) ===")
    cmd_status(b)
    mem = ROOT / ".claude" / "memory" / "MEMORY.md"
    if mem.exists():
        print("\n=== Mémoire projet (.claude/memory/MEMORY.md) ===")
        print(mem.read_text(encoding="utf-8")[:3000])
    j = ROOT / ".claude" / "memory" / "journal.md"
    if j.exists():
        print("\n=== Journal (dernières lignes) ===")
        print("\n".join(j.read_text(encoding="utf-8").splitlines()[-12:]))
    print("\nRappel : lire CLAUDE.md ; ne jamais passer une tâche à 'Validé' ; s'arrêter à la gate de phase.")


def cmd_hook_stop(b):
    """Pousse Claude à finir la phase en mode autopilot (max 5 relances)."""
    if not (STATE / "autopilot").exists():
        return
    n_file = STATE / "stop_nudges"
    n = int(n_file.read_text() or 0) if n_file.exists() else 0
    t = b.next_task()
    if t is None or n >= 5:
        return
    n_file.write_text(str(n + 1))
    print(json.dumps({"decision": "block", "reason":
        f"Mode autopilot : la tâche {t['ID']} ({t['Statut']}) n'est pas terminée. Continue la phase en cours, "
        "ou si tu es réellement bloqué, note-le dans la colonne 'Preuve / notes', ajoute une hypothèse H-xx et passe à la suite."}))


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("status", "next", "export-md", "hook-session", "hook-stop"):
        sub.add_parser(name)
    sub.add_parser("show").add_argument("id")
    s = sub.add_parser("set")
    s.add_argument("id")
    s.add_argument("statut")
    s.add_argument("--note", default="")
    sub.add_parser("validate").add_argument("phase")
    sub.add_parser("hypo").add_argument("ids", nargs="*")
    ha = sub.add_parser("hypo-add")
    for k in ("tasks", "question", "assumption", "impact", "ask"):
        ha.add_argument(f"--{k}", required=True)
    a = ap.parse_args()
    b = Backlog()
    if a.cmd == "status":
        cmd_status(b)
    elif a.cmd == "next":
        t = b.next_task()
        print(fmt(t) if t else "Aucune tâche à exécuter (phase à valider ou projet terminé).")
    elif a.cmd == "show":
        print(fmt(b.task(a.id)))
    elif a.cmd == "set":
        cmd_set(b, a.id, a.statut, a.note)
    elif a.cmd == "validate":
        cmd_validate(b, a.phase)
    elif a.cmd == "hypo":
        cmd_hypo(b, a.ids)
    elif a.cmd == "hypo-add":
        cmd_hypo_add(b, a)
    elif a.cmd == "export-md":
        cmd_export_md(b)
    elif a.cmd == "hook-session":
        cmd_hook_session(b)
    elif a.cmd == "hook-stop":
        cmd_hook_stop(b)


if __name__ == "__main__":
    main()
